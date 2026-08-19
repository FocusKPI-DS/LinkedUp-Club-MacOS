#!/usr/bin/env python3
"""Generate Xcode App Store Upload Guide (Word document, English only, no emoji)."""

import os
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
IMAGES_DIR = os.path.join(SCRIPT_DIR, "images")
OUTPUT_PATH = os.path.join(SCRIPT_DIR, "Xcode_App_Store_Upload_Guide.docx")


def find_image(name):
    for ext in ["png", "jpg", "jpeg", "webp", "PNG", "JPG"]:
        path = os.path.join(IMAGES_DIR, f"{name}.{ext}")
        if os.path.exists(path):
            return path
    return None


def set_cell_shading(cell, color_hex):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = tc_pr.makeelement(qn("w:shd"), {
        qn("w:val"): "clear",
        qn("w:color"): "auto",
        qn("w:fill"): color_hex,
    })
    tc_pr.append(shd)


def add_styled_table(doc, headers, rows, header_color="1F4E79"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(10.5)
        set_cell_shading(cell, header_color)
    for r, row_data in enumerate(rows):
        for c, cell_text in enumerate(row_data):
            cell = table.rows[r + 1].cells[c]
            cell.text = cell_text
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10.5)
            if r % 2 == 1:
                set_cell_shading(cell, "F2F2F2")
    return table


def add_callout(doc, text, label="Note", color=(0x0B, 0x5A, 0x9D)):
    p = doc.add_paragraph()
    run_label = p.add_run(f"{label}: ")
    run_label.bold = True
    run_label.font.size = Pt(10.5)
    run_label.font.color.rgb = RGBColor(*color)
    run_text = p.add_run(text)
    run_text.font.size = Pt(10.5)


def try_add_image(doc, image_name, caption, width=Inches(5.5)):
    img_path = find_image(image_name)
    if img_path:
        doc.add_picture(img_path, width=width)
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        run.italic = True
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"[{caption}] -- Screenshot to be inserted")
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        run.italic = True


def add_body(doc, text):
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.size = Pt(11)
    return p


def add_numbered_list(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Number")
        for run in p.runs:
            run.font.size = Pt(11)


def add_bullet_list(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            run.font.size = Pt(11)


def build_document():
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ── Cover ──
    title = doc.add_heading("Xcode App Store Upload Guide", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta_info = [
        ("Audience", "New Team Members"),
        ("Platforms", "iOS and macOS"),
        ("Developer Account", "Focus KPI"),
        ("Last Updated", "August 18, 2026"),
    ]
    for label, value in meta_info:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_l = p.add_run(f"{label}: ")
        run_l.bold = True
        run_l.font.size = Pt(11)
        run_v = p.add_run(value)
        run_v.font.size = Pt(11)

    doc.add_page_break()

    # ── Table of Contents ──
    doc.add_heading("Table of Contents", level=1)
    toc_items = [
        "Step 1: Update the Version Number",
        "Step 2: Configure Xcode and Build the Project",
        "Step 3: Archive the Build",
        "Step 4: Distribute to the App Store",
        "Step 5: Submit for Review on App Store Connect",
        "Frequently Asked Questions",
    ]
    for i, item in enumerate(toc_items, 1):
        p = doc.add_paragraph(f"{i}. {item}")
        for run in p.runs:
            run.font.size = Pt(11)

    doc.add_page_break()

    # ══════════════════════════════════════════════
    # Step 1
    # ══════════════════════════════════════════════
    doc.add_heading("Step 1: Update the Version Number", level=1)

    add_callout(doc,
        "Make sure the new version has been fully tested and is free of critical bugs before starting the release process.",
        label="Important", color=(0xCC, 0x00, 0x00))

    add_body(doc,
        "Once the new version has passed all tests, instruct your AI Agent "
        "(Cursor, Antigravity, Claude Code, Codex, or any other available tool) "
        "to update the version number."
    )

    doc.add_heading("Key Requirements", level=2)
    add_bullet_list(doc, [
        "iOS and macOS version numbers must be updated separately. Do not share the same configuration between them.",
        "Have the AI confirm that CFBundleShortVersionString (version number) and CFBundleVersion (build number) "
        "in ios/Runner/Info.plist have been correctly incremented.",
        "Have the AI confirm that CFBundleShortVersionString (version number) and CFBundleVersion (build number) "
        "in macos/Runner/Info.plist have been correctly incremented.",
    ])

    doc.add_heading("Example Prompt", level=2)
    p = doc.add_paragraph()
    run = p.add_run(
        "Please update the iOS version from 2.1 to 2.2 and the build number from 70 to 71.\n"
        "Also update the macOS version from 2.0.0 to 2.1.0 and the build number from 71 to 72."
    )
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    p.paragraph_format.left_indent = Cm(1)

    # ══════════════════════════════════════════════
    # Step 2
    # ══════════════════════════════════════════════
    doc.add_heading("Step 2: Configure Xcode and Build the Project", level=1)

    doc.add_heading("2.1 Open the Correct Xcode Project", level=2)
    add_body(doc,
        "Open the Xcode project files for iOS and macOS separately. "
        "You can verify which platform you are working on by checking the toolbar "
        "at the top of the Xcode window."
    )

    doc.add_heading("2.2 Select the Correct Build Destination", level=2)
    add_callout(doc,
        "This step is critical. You must select the correct target device, otherwise "
        "the Archive option will be grayed out.",
        label="Warning", color=(0xCC, 0x66, 0x00))

    add_styled_table(doc,
        headers=["Platform", "Build Destination"],
        rows=[
            ["macOS", "Any Mac (arm64, x86_64)"],
            ["iOS", "Any iOS Device"],
        ],
    )

    doc.add_paragraph()
    try_add_image(doc, "step2_select_any_mac",
                  "Figure 1: Selecting Any Mac as the Build Destination")

    doc.add_heading("2.3 Handling Build Warnings and Errors", level=2)
    add_body(doc,
        "During the build process, the left-hand navigator may display various issues:"
    )

    add_styled_table(doc,
        headers=["Type", "Icon", "Action Required"],
        rows=[
            ["Yellow Warning", "Yellow triangle", "Can be safely ignored; does not affect archiving"],
            ["Red Error", "Red circle", "Must be resolved -- take a screenshot and submit it to your AI Agent"],
        ],
    )

    doc.add_heading("2.4 Wait for the Build to Complete", level=2)
    add_body(doc,
        'When the toolbar displays "Build Succeeded", the build is complete and you '
        "can proceed to the next step."
    )

    # ══════════════════════════════════════════════
    # Step 3
    # ══════════════════════════════════════════════
    doc.add_heading("Step 3: Archive the Build", level=1)

    doc.add_heading("3.1 Create the Archive", level=2)
    add_numbered_list(doc, [
        'Confirm that the toolbar shows "Build Succeeded".',
        "In the Xcode menu bar at the top of the screen, click Product.",
        "From the dropdown menu, select Archive.",
    ])

    try_add_image(doc, "step3_product_archive",
                  "Figure 2: Selecting Product > Archive from the menu bar")

    doc.add_heading("3.2 Wait for Archiving to Finish", level=2)
    add_bullet_list(doc, [
        "The archiving process takes approximately 5 minutes.",
        "Do not close Xcode or perform other operations while archiving is in progress.",
        "When archiving is complete, the Organizer window will appear automatically, "
        "displaying a list of all available archives.",
    ])

    try_add_image(doc, "step3_organizer_archives",
                  "Figure 3: The Organizer window showing completed archives")

    # ══════════════════════════════════════════════
    # Step 4
    # ══════════════════════════════════════════════
    doc.add_heading("Step 4: Distribute to the App Store", level=1)

    doc.add_heading("4.1 Start Distribution", level=2)
    add_numbered_list(doc, [
        "In the Organizer window, select the most recent archive "
        "(verify by checking the Creation Date column).",
        'Click the "Distribute App" button on the right side of the window.',
    ])

    doc.add_heading("4.2 Choose a Distribution Method", level=2)
    add_body(doc, "In the distribution method dialog, choose based on your needs:")

    add_styled_table(doc,
        headers=["Method", "Purpose", "Description"],
        rows=[
            ["App Store Connect", "Production Release",
             "Publish to the App Store for end users to download"],
            ["TestFlight Internal Only", "Internal Testing",
             "Available only to internal testers"],
            ["Direct Distribution", "Direct Distribution",
             "Distribute outside the App Store"],
            ["Debugging", "Debug Builds",
             "For debugging purposes only"],
        ],
    )

    doc.add_paragraph()
    add_callout(doc,
        "Selecting App Store Connect allows you to use the build for both "
        "production release and TestFlight testing. If you only want to release "
        "a test version, you may also choose TestFlight Internal Only.",
        label="Tip")

    try_add_image(doc, "step4_distribute_method",
                  "Figure 4: Selecting a distribution method")

    add_body(doc,
        'After making your selection, click the "Distribute" button '
        "in the lower-right corner to begin uploading."
    )

    doc.add_heading("4.3 Interpreting Upload Results", level=2)
    add_styled_table(doc,
        headers=["Result", "Explanation"],
        rows=[
            ["Upload succeeded with yellow warning",
             "This is normal. The upload was successful; yellow warnings can be ignored."],
            ["Upload failed with red error",
             "Action required -- take a screenshot and submit it to your AI Agent for troubleshooting."],
        ],
    )

    # ══════════════════════════════════════════════
    # Step 5
    # ══════════════════════════════════════════════
    doc.add_heading("Step 5: Submit for Review on App Store Connect", level=1)

    doc.add_heading("5.1 Log In to App Store Connect", level=2)
    add_numbered_list(doc, [
        "Open your browser and navigate to https://appstoreconnect.apple.com",
        "Sign in with the Focus KPI developer account.",
    ])

    doc.add_heading("5.2 Wait for Build Processing", level=2)
    add_callout(doc,
        "After a successful Distribute from Xcode, allow approximately "
        "20 to 30 minutes for Apple's servers to finish processing the uploaded build.",
        label="Note")

    doc.add_heading("5.3 Select the Build and Submit for Review", level=2)
    add_body(doc, "Using macOS as an example:")

    add_numbered_list(doc, [
        'In the left sidebar, locate the version under macOS App (e.g., "2.00 Prepare for Submission").',
        "In the Build section in the center of the page, find and select the build you just uploaded.",
        'Fill in the release notes under "What\'s New in This Version".',
        'Click "Save" in the upper-right corner.',
        'Click "Add for Review" to submit the build for Apple\'s review.',
    ])

    try_add_image(doc, "step5_app_store_connect",
                  "Figure 5: The App Store Connect submission page")

    doc.add_heading("5.4 Wait for Review Approval", level=2)
    add_callout(doc,
        "After submission, Apple's review team typically takes about 2 days to "
        "complete the review. Once approved, you will receive a notification email "
        "and the new version will go live on the App Store.",
        label="Important", color=(0xCC, 0x00, 0x00))

    # ══════════════════════════════════════════════
    # FAQ
    # ══════════════════════════════════════════════
    doc.add_page_break()
    doc.add_heading("Frequently Asked Questions", level=1)

    qa_list = [
        (
            "Q1: Can I archive and upload iOS and macOS at the same time?",
            "Yes, but it is recommended to complete one platform at a time to avoid confusion.",
        ),
        (
            "Q2: The Archive option in the Product menu is grayed out. Why?",
            "Make sure you have selected the correct Build Destination. "
            "For macOS projects, select Any Mac. For iOS projects, select Any iOS Device. "
            "Do not select a Simulator.",
        ),
        (
            "Q3: I am getting a code signing error during upload.",
            "Take a screenshot and send it to your AI Agent. Common causes include "
            "expired certificates, mismatched Provisioning Profiles, or insufficient "
            "developer account permissions.",
        ),
        (
            "Q4: I cannot see my uploaded build in App Store Connect.",
            "Confirm that the Distribute step completed successfully with no red errors. "
            "Allow 20 to 30 minutes for Apple's servers to process the build. "
            "Also check your registered email, as Apple may send a notification if processing fails.",
        ),
        (
            "Q5: My submission was rejected by Apple. What should I do?",
            "Review the rejection reason provided by Apple, fix the issue, re-upload, "
            "and submit again. You can share the rejection details with your AI Agent for analysis.",
        ),
    ]

    for question, answer in qa_list:
        doc.add_heading(question, level=2)
        add_body(doc, answer)

    # ══════════════════════════════════════════════
    # Summary
    # ══════════════════════════════════════════════
    doc.add_heading("Process Overview", level=1)
    add_body(doc,
        "Testing complete -> AI updates version number -> Xcode Build -> "
        "Archive (~5 min) -> Distribute upload -> App Store Connect select build "
        "(wait ~30 min) -> Fill in release notes -> Add for Review -> "
        "Wait for review (~2 days) -> Published on the App Store"
    )

    doc.add_paragraph()
    add_callout(doc,
        "If you encounter any red errors at any point in this process, take a "
        "screenshot immediately and send it to your AI Agent (Cursor, Antigravity, "
        "Claude Code, Codex, etc.). They can quickly help you diagnose and resolve "
        "most build and upload issues.",
        label="Tip")

    doc.add_paragraph()
    add_callout(doc,
        "If you need to deploy the web version, there is an existing deploy script "
        "in the project repository. Simply ask your AI Agent to locate and run it.",
        label="Note")

    # ── Save ──
    doc.save(OUTPUT_PATH)
    print(f"Done. Word document saved to: {OUTPUT_PATH}")


if __name__ == "__main__":
    build_document()
