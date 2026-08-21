#!/usr/bin/env python3
"""Generate Lona Project Credentials Document (Word, English only, no emoji)."""

import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_PATH = os.path.join(SCRIPT_DIR, "Lona_Project_Credentials.docx")


def set_cell_shading(cell, color_hex):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = tc_pr.makeelement(qn("w:shd"), {
        qn("w:val"): "clear",
        qn("w:color"): "auto",
        qn("w:fill"): color_hex,
    })
    tc_pr.append(shd)


def add_credentials_table(doc, headers, rows, header_color="1F4E79"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(10.5)
        set_cell_shading(cell, header_color)
    for r, row_data in enumerate(rows):
        for c, cell_text in enumerate(row_data):
            cell = table.rows[r + 1].cells[c]
            cell.text = cell_text
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10.5)
                    if headers[c] == "Password":
                        run.font.name = "Courier New"
            if r % 2 == 1:
                set_cell_shading(cell, "F2F2F2")
    return table


def add_callout(doc, text, label="Note", color=(0x0B, 0x5A, 0x9D)):
    p = doc.add_paragraph()
    run_label = p.add_run(f"{label}: ")
    run_label.bold = True
    run_label.font.size = Pt(10.5)
    run_label.font.color.rgb = RGBColor(*color)
    run_text = p.add_run(text)
    run_text.font.size = Pt(10.5)


def add_body(doc, text):
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.size = Pt(11)
    return p


def build_document():
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ── Cover ──
    title = doc.add_heading("Lona Project Credentials", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("CONFIDENTIAL -- Internal Use Only")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)

    meta_info = [
        ("Project", "Lona (LinkedUp Club)"),
        ("Maintained By", "Focus KPI Development Team"),
        ("Last Updated", "August 18, 2026"),
    ]
    doc.add_paragraph()
    for label, value in meta_info:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_l = p.add_run(f"{label}: ")
        run_l.bold = True
        run_l.font.size = Pt(11)
        run_v = p.add_run(value)
        run_v.font.size = Pt(11)

    doc.add_paragraph()
    add_callout(doc,
        "This document contains sensitive credentials. Do not share it outside "
        "the team. Do not commit it to any repository. Do not upload it to any "
        "cloud storage without proper access controls. Store it securely at all times.",
        label="Warning", color=(0xCC, 0x00, 0x00))

    doc.add_page_break()

    # ══════════════════════════════════════════════
    # Apple Developer / App Store Connect
    # ══════════════════════════════════════════════
    doc.add_heading("1. Apple Developer / App Store Connect", level=1)

    add_body(doc,
        "Used for publishing iOS and macOS apps to the App Store, managing "
        "TestFlight builds, and accessing App Store Connect."
    )

    add_credentials_table(doc,
        headers=["Field", "Value"],
        rows=[
            ["Service", "App Store Connect"],
            ["URL", "https://appstoreconnect.apple.com"],
            ["Account Email", "appdev@focuskpi.com"],
            ["Password", "LecQ#x!rIUl9"],
            ["Team Name", "Dan Zhang"],
            ["Team ID", "3GQ93ABMNG"],
            ["Bundle ID (macOS)", "com.focuskpi.linkedup"],
        ],
    )

    doc.add_paragraph()
    add_callout(doc,
        "This account is also used to log in to the Apple Developer portal "
        "at https://developer.apple.com for managing certificates, "
        "provisioning profiles, and app identifiers.",
        label="Note")

    # ══════════════════════════════════════════════
    # Placeholder sections for future credentials
    # ══════════════════════════════════════════════
    doc.add_heading("2. Firebase", level=1)
    add_body(doc, "Firebase project used for backend services, authentication, and hosting.")
    add_credentials_table(doc,
        headers=["Field", "Value"],
        rows=[
            ["Service", "Firebase Console"],
            ["URL", "https://console.firebase.google.com"],
            ["Project ID", "linkedup-c3e29"],
            ["Account Email", "(to be added)"],
            ["Password", "(to be added)"],
        ],
    )

    doc.add_paragraph()

    doc.add_heading("3. GitHub", level=1)
    add_body(doc, "Source code repository hosting.")
    add_credentials_table(doc,
        headers=["Field", "Value"],
        rows=[
            ["Service", "GitHub"],
            ["URL", "https://github.com/FocusKPI-DS"],
            ["Organization", "FocusKPI-DS"],
            ["Repository", "LinkedUp-Club-MacOS"],
            ["Account Email", "(to be added)"],
            ["Password / Token", "(to be added)"],
        ],
    )

    doc.add_paragraph()

    doc.add_heading("4. Google OAuth (Desktop App)", level=1)
    add_body(doc, "OAuth credentials for Google Sign-In on macOS and iOS.")
    add_credentials_table(doc,
        headers=["Field", "Value"],
        rows=[
            ["Service", "Google Cloud Console"],
            ["URL", "https://console.cloud.google.com"],
            ["Client ID", "548534727055-042f43gm4l2jhs3qf4t69m4g4g6mnjjc.apps.googleusercontent.com"],
            ["Client Secret", "(see env.json)"],
        ],
    )

    doc.add_paragraph()

    doc.add_heading("5. Additional Accounts", level=1)
    add_body(doc,
        "Add any additional service credentials below as the project grows. "
        "Use the same table format for consistency."
    )
    add_credentials_table(doc,
        headers=["Service", "URL", "Account Email", "Password", "Notes"],
        rows=[
            ["(service name)", "(login URL)", "(email)", "(password)", "(notes)"],
        ],
    )

    # ══════════════════════════════════════════════
    # Change Log
    # ══════════════════════════════════════════════
    doc.add_page_break()
    doc.add_heading("Change Log", level=1)
    add_body(doc, "Track all credential updates below for audit purposes.")

    add_credentials_table(doc,
        headers=["Date", "Changed By", "Description"],
        rows=[
            ["2026-08-18", "Initial Setup", "Added Apple Developer / App Store Connect credentials"],
        ],
    )

    # ── Save ──
    doc.save(OUTPUT_PATH)
    print(f"Done. Word document saved to: {OUTPUT_PATH}")


if __name__ == "__main__":
    build_document()
