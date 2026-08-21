#!/usr/bin/env python3
"""Generate Lona Local Development Setup Guide (Word document, English only, no emoji)."""

import os
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_PATH = os.path.join(SCRIPT_DIR, "Lona_Local_Dev_Setup_Guide.docx")


def set_cell_shading(cell, color_hex):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = tc_pr.makeelement(qn("w:shd"), {
        qn("w:val"): "clear",
        qn("w:color"): "auto",
        qn("w:fill"): color_hex,
    })
    tc_pr.append(shd)


def add_styled_table(doc, headers, rows, header_color="1F4E79"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(10.5)
        set_cell_shading(cell, header_color)
    for r, row_data in enumerate(rows):
        for c, cell_text in enumerate(row_data):
            cell = table.rows[r + 1].cells[c]
            cell.text = cell_text
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10.5)
            if r % 2 == 1:
                set_cell_shading(cell, "F2F2F2")
    return table


def add_callout(doc, text, label="Note", color=(0x0B, 0x5A, 0x9D)):
    p = doc.add_paragraph()
    run_label = p.add_run(f"{label}: ")
    run_label.bold = True
    run_label.font.size = Pt(10.5)
    run_label.font.color.rgb = RGBColor(*color)
    run_text = p.add_run(text)
    run_text.font.size = Pt(10.5)


def add_body(doc, text):
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.size = Pt(11)
    return p


def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    run = p.add_run(code_text)
    run.font.name = "Courier New"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_numbered_list(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Number")
        for run in p.runs:
            run.font.size = Pt(11)


def add_bullet_list(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            run.font.size = Pt(11)


def build_document():
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ── Cover ──
    title = doc.add_heading("Lona Local Development Setup Guide", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta_info = [
        ("Audience", "New Team Members"),
        ("Project", "Lona (LinkedUp Club)"),
        ("Framework", "Flutter 3.38.x / Dart 3.10.x"),
        ("Platforms", "macOS, iOS, Web"),
        ("Last Updated", "August 18, 2026"),
    ]
    for label, value in meta_info:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_l = p.add_run(f"{label}: ")
        run_l.bold = True
        run_l.font.size = Pt(11)
        run_v = p.add_run(value)
        run_v.font.size = Pt(11)

    doc.add_page_break()

    # ── Table of Contents ──
    doc.add_heading("Table of Contents", level=1)
    toc_items = [
        "Step 1: Prerequisites",
        "Step 2: Clone the Repository",
        "Step 3: Environment Configuration",
        "Step 4: Install Dependencies (flutter pub get)",
        "Step 5: Run the App Locally (flutter run)",
        "Step 6: Platform-Specific Notes",
        "Common Issues and Troubleshooting",
    ]
    for i, item in enumerate(toc_items, 1):
        p = doc.add_paragraph(f"{i}. {item}")
        for run in p.runs:
            run.font.size = Pt(11)

    doc.add_page_break()

    # ══════════════════════════════════════════════
    # Step 1: Prerequisites
    # ══════════════════════════════════════════════
    doc.add_heading("Step 1: Prerequisites", level=1)

    add_body(doc,
        "Before setting up the Lona project locally, make sure the following "
        "tools are installed on your Mac."
    )

    add_styled_table(doc,
        headers=["Tool", "Required Version", "How to Install"],
        rows=[
            ["Flutter SDK", "3.38.x (stable channel)",
             "https://docs.flutter.dev/get-started/install/macos"],
            ["Dart SDK", "3.10.x (bundled with Flutter)", "Included with Flutter"],
            ["Xcode", "Latest version from Mac App Store",
             "Mac App Store or developer.apple.com"],
            ["Xcode Command Line Tools", "Latest",
             "Run: xcode-select --install"],
            ["CocoaPods", "Latest",
             "Run: sudo gem install cocoapods"],
            ["Git", "Latest", "Bundled with Xcode CLI tools"],
            ["Node.js and npm", "LTS version (for web/landing page)",
             "https://nodejs.org"],
            ["Firebase CLI", "Latest (for web deployment)",
             "Run: npm install -g firebase-tools"],
        ],
    )

    doc.add_paragraph()
    add_callout(doc,
        "Run 'flutter doctor' in your terminal to verify that all required "
        "tools are properly installed and configured. Fix any issues it reports "
        "before proceeding.",
        label="Tip")

    doc.add_heading("Verify Flutter Installation", level=2)
    add_body(doc, "Run the following command to check your Flutter setup:")
    add_code_block(doc, "flutter doctor -v")

    add_body(doc,
        "Ensure there are no red errors. Yellow warnings for unused platforms "
        "(e.g., Android, Linux) can be ignored if you are only targeting macOS, iOS, and Web."
    )

    # ══════════════════════════════════════════════
    # Step 2: Clone the Repository
    # ══════════════════════════════════════════════
    doc.add_heading("Step 2: Clone the Repository", level=1)

    add_body(doc, "Clone the project from the GitHub repository:")
    add_code_block(doc,
        "git clone https://github.com/FocusKPI-DS/LinkedUp-Club-MacOS.git\n"
        "cd LinkedUp-Club-MacOS"
    )

    add_callout(doc,
        "Make sure you have access to the FocusKPI-DS GitHub organization. "
        "If you do not have access, ask your team lead to add you.",
        label="Note")

    # ══════════════════════════════════════════════
    # Step 3: Environment Configuration
    # ══════════════════════════════════════════════
    doc.add_heading("Step 3: Environment Configuration", level=1)

    add_body(doc,
        "The project requires an environment configuration file for Google OAuth. "
        "Create this file before running the app."
    )

    doc.add_heading("3.1 Create the env.json File", level=2)
    add_numbered_list(doc, [
        "In the project root directory, find the file env.json.example.",
        "Copy it and rename the copy to env.json.",
        "Fill in the required values (ask your team lead for the credentials if needed).",
    ])
    add_code_block(doc, "cp env.json.example env.json")

    add_body(doc, "The env.json file should contain:")
    add_code_block(doc,
        '{\n'
        '  "GOOGLE_DESKTOP_OAUTH_CLIENT_ID": "<your-client-id>",\n'
        '  "GOOGLE_DESKTOP_OAUTH_CLIENT_SECRET": "<your-client-secret>"\n'
        '}'
    )

    add_callout(doc,
        "Never commit env.json to the repository. It is already included in .gitignore.",
        label="Warning", color=(0xCC, 0x66, 0x00))

    doc.add_heading("3.2 Firebase Configuration", level=2)
    add_body(doc,
        "The project uses Firebase for backend services. The Firebase configuration "
        "files (GoogleService-Info.plist, firebase_options.dart) are already included "
        "in the repository. No additional Firebase setup is required for local development."
    )

    # ══════════════════════════════════════════════
    # Step 4: Install Dependencies
    # ══════════════════════════════════════════════
    doc.add_heading("Step 4: Install Dependencies (flutter pub get)", level=1)

    add_body(doc,
        "This step downloads all Dart and Flutter packages listed in pubspec.yaml, "
        "including local dependencies in the dependencies/ folder."
    )

    doc.add_heading("4.1 Run flutter pub get", level=2)
    add_code_block(doc, "flutter pub get")

    add_body(doc,
        "This command resolves and downloads all dependencies. It typically takes "
        "1 to 2 minutes on first run. Subsequent runs are faster due to caching."
    )

    doc.add_heading("4.2 Install iOS CocoaPods Dependencies", level=2)
    add_body(doc, "If you plan to run on iOS, also install the CocoaPods dependencies:")
    add_code_block(doc,
        "cd ios\npod install\ncd .."
    )

    doc.add_heading("4.3 Install macOS CocoaPods Dependencies", level=2)
    add_body(doc, "If you plan to run on macOS, install CocoaPods for the macOS target:")
    add_code_block(doc,
        "cd macos\npod install\ncd .."
    )

    add_callout(doc,
        "If 'pod install' fails, try running 'pod repo update' first, then retry. "
        "You can also try 'pod install --repo-update' to combine both steps.",
        label="Tip")

    doc.add_heading("4.4 Local Dependencies", level=2)
    add_body(doc,
        "The project includes three local packages in the dependencies/ directory. "
        "These are resolved automatically by flutter pub get:"
    )
    add_styled_table(doc,
        headers=["Package", "Path"],
        rows=[
            ["ff_commons", "dependencies/ff_commons"],
            ["ff_theme", "dependencies/ff_theme"],
            ["branchio_dynamic_linking_akp5u6", "dependencies/branchio_dynamic_linking_akp5u6"],
        ],
    )

    # ══════════════════════════════════════════════
    # Step 5: Run the App Locally
    # ══════════════════════════════════════════════
    doc.add_heading("Step 5: Run the App Locally (flutter run)", level=1)

    add_body(doc,
        "Once dependencies are installed, you can run the app on your target platform."
    )

    doc.add_heading("5.1 Run on macOS", level=2)
    add_code_block(doc, "flutter run -d macos")
    add_body(doc,
        "This launches the Lona app as a native macOS application. "
        "The first build may take several minutes as it compiles all native dependencies."
    )

    doc.add_heading("5.2 Run on iOS Simulator", level=2)
    add_code_block(doc, "flutter run -d ios")
    add_body(doc,
        "This launches the app in the iOS Simulator. Make sure you have a simulator "
        "set up in Xcode (Xcode > Settings > Platforms > Download a simulator if needed)."
    )
    add_body(doc, "To list available devices:")
    add_code_block(doc, "flutter devices")

    doc.add_heading("5.3 Run on Chrome (Web)", level=2)
    add_code_block(doc, "flutter run -d chrome")
    add_body(doc,
        "This launches the app in a Chrome browser window. Web mode is useful for "
        "quick iteration but may have some feature differences compared to native platforms."
    )

    doc.add_heading("5.4 Run from Xcode (Alternative)", level=2)
    add_body(doc,
        "You can also open the Xcode project directly and run from there. "
        "This is especially useful for debugging native platform issues."
    )
    add_numbered_list(doc, [
        "For macOS: Open macos/Runner.xcworkspace in Xcode.",
        "For iOS: Open ios/Runner.xcworkspace in Xcode.",
        "Select the appropriate scheme and target device.",
        "Click the Run button or press Cmd+R.",
    ])

    add_callout(doc,
        "Always open the .xcworkspace file, not the .xcodeproj file. "
        "The workspace includes the CocoaPods dependencies.",
        label="Important", color=(0xCC, 0x00, 0x00))

    doc.add_heading("5.5 Hot Reload and Hot Restart", level=2)
    add_body(doc,
        "When running via flutter run in the terminal, you can use the "
        "following keyboard shortcuts:"
    )
    add_styled_table(doc,
        headers=["Key", "Action", "Description"],
        rows=[
            ["r", "Hot Reload", "Applies code changes without restarting the app. Preserves app state."],
            ["R", "Hot Restart", "Restarts the app from scratch. Resets all state."],
            ["q", "Quit", "Stops the running app."],
            ["h", "Help", "Shows all available commands."],
        ],
    )

    # ══════════════════════════════════════════════
    # Step 6: Platform-Specific Notes
    # ══════════════════════════════════════════════
    doc.add_heading("Step 6: Platform-Specific Notes", level=1)

    doc.add_heading("macOS Entitlements", level=2)
    add_body(doc,
        "The macOS build requires specific entitlements for features such as "
        "network access, camera, microphone, and file system access. These are "
        "pre-configured in the macos/Runner/*.entitlements files. Do not modify "
        "them unless you know what you are doing."
    )

    doc.add_heading("iOS Signing", level=2)
    add_body(doc,
        "For running on a physical iOS device, you need to configure code signing "
        "in Xcode. Open ios/Runner.xcworkspace, go to Signing & Capabilities, "
        "and select the Focus KPI team. For Simulator testing, automatic signing "
        "with your personal Apple ID is sufficient."
    )

    doc.add_heading("Web Build for Deployment", level=2)
    add_body(doc,
        "To build and deploy the web version, use the existing deploy script "
        "in the project root:"
    )
    add_code_block(doc, "./deploy.sh")
    add_body(doc,
        "This script builds the Next.js landing page, builds the Flutter web app, "
        "combines them, and deploys to Firebase Hosting. The deployed URLs are:"
    )
    add_bullet_list(doc, [
        "Landing page: https://linkedup-c3e29.web.app",
        "Flutter app: https://linkedup-c3e29.web.app/app",
    ])

    # ══════════════════════════════════════════════
    # FAQ
    # ══════════════════════════════════════════════
    doc.add_page_break()
    doc.add_heading("Common Issues and Troubleshooting", level=1)

    qa_list = [
        (
            'Q1: "flutter pub get" fails with dependency conflicts',
            "Try running 'flutter clean' first, then 'flutter pub get' again. "
            "If the issue persists, delete the pubspec.lock file and retry. "
            "You can also ask your AI Agent to analyze the error output."
        ),
        (
            'Q2: CocoaPods errors during pod install',
            "Common fixes:\n"
            "- Run 'pod repo update' to refresh the local CocoaPods spec repository.\n"
            "- Delete the Podfile.lock and Pods/ directory, then run 'pod install' again.\n"
            "- Make sure your CocoaPods version is up to date: 'sudo gem install cocoapods'."
        ),
        (
            'Q3: Xcode build fails with signing errors',
            "Open the .xcworkspace in Xcode, go to the Runner target > Signing & Capabilities, "
            "and make sure the correct team is selected. For development, automatic signing "
            "is usually sufficient."
        ),
        (
            'Q4: The app runs but shows a blank or white screen',
            "This usually indicates a Firebase configuration issue. Make sure the "
            "GoogleService-Info.plist files are in the correct locations:\n"
            "- iOS: ios/Runner/GoogleService-Info.plist\n"
            "- macOS: macos/Runner/GoogleService-Info.plist"
        ),
        (
            'Q5: flutter run -d macos is very slow on first build',
            "The first macOS build compiles all native dependencies and can take "
            "5 to 10 minutes. Subsequent builds are much faster due to incremental "
            "compilation. Using 'flutter run' (hot reload) instead of stopping and "
            "restarting will save significant time during development."
        ),
        (
            'Q6: "No devices found" when running flutter run',
            "Run 'flutter devices' to see available targets. For macOS, ensure you "
            "have enabled macOS desktop support: 'flutter config --enable-macos-desktop'. "
            "For iOS Simulator, open Xcode and start a simulator first."
        ),
        (
            'Q7: How do I switch between Flutter channels?',
            "The project uses the stable channel. Verify with 'flutter channel'. "
            "If you are on a different channel, switch with 'flutter channel stable' "
            "followed by 'flutter upgrade'."
        ),
    ]

    for question, answer in qa_list:
        doc.add_heading(question, level=2)
        add_body(doc, answer)

    # ══════════════════════════════════════════════
    # Quick Reference
    # ══════════════════════════════════════════════
    doc.add_heading("Quick Reference: Common Commands", level=1)

    add_styled_table(doc,
        headers=["Command", "Purpose"],
        rows=[
            ["flutter doctor -v", "Check development environment setup"],
            ["flutter pub get", "Install/update all Dart dependencies"],
            ["flutter clean", "Clean build artifacts (useful for fixing build issues)"],
            ["flutter run -d macos", "Run the app on macOS"],
            ["flutter run -d ios", "Run the app on iOS Simulator"],
            ["flutter run -d chrome", "Run the app in Chrome browser"],
            ["flutter devices", "List all available target devices"],
            ["flutter build macos --release", "Build a release version for macOS"],
            ["flutter build ios --release", "Build a release version for iOS"],
            ["flutter build web --release", "Build a release version for Web"],
            ["./deploy.sh", "Build and deploy the web version to Firebase"],
            ["./build_macos.sh", "Build macOS app and create DMG installer"],
        ],
    )

    doc.add_paragraph()
    add_callout(doc,
        "If you encounter any errors during setup or development, take a screenshot "
        "or copy the terminal output and send it to your AI Agent (Cursor, Antigravity, "
        "Claude Code, Codex, etc.) for assistance.",
        label="Tip")

    # ── Save ──
    doc.save(OUTPUT_PATH)
    print(f"Done. Word document saved to: {OUTPUT_PATH}")


if __name__ == "__main__":
    build_document()
